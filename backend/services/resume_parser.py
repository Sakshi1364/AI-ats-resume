import io
from typing import Tuple, Optional

import pdfplumber
from docx import Document
import PyPDF2

from backend.utils.file_utils import (
    TextExtractionError,
    log_error,
    log_warning,
    log_info,
    with_fallback
)

from backend.core.config import (
    MAX_FILE_SIZE_BYTES,
    MAX_FILE_SIZE_MB
)


class FileParsingError(Exception):
    pass


class FileValidationError(Exception):
    pass


def validate_file(
    file_data: bytes,
    filename: str
) -> Tuple[bool, str, Optional[str]]:

    # -----------------------------
    # 1. Check file size
    # -----------------------------
    file_size_bytes = len(file_data)

    if file_size_bytes > MAX_FILE_SIZE_BYTES:
        size_mb = file_size_bytes / (1024 * 1024)

        return False, (
            f"File size ({size_mb:.2f} MB) exceeds the maximum of "
            f"{MAX_FILE_SIZE_MB} MB. "
            "Please upload a smaller file or compress your resume."
        ), None

    # -----------------------------
    # 2. Check empty file
    # -----------------------------
    if file_size_bytes == 0:
        return False, (
            "Uploaded file is empty. "
            "Please check the file you have uploaded."
        ), None

    # -----------------------------
    # 3. Detect file type
    # -----------------------------
    filename_lower = filename.lower()

    # PDF
    if filename_lower.endswith(".pdf"):

        if file_data.startswith(b"%PDF"):
            return True, "", "pdf"

        return False, (
            "The uploaded PDF file appears to be invalid or corrupted."
        ), None

    # DOCX
    if filename_lower.endswith(".docx"):

        # DOCX files are ZIP-based and start with PK
        if file_data.startswith(b"PK"):
            return True, "", "docx"

        return False, (
            "The uploaded DOCX file appears to be invalid or corrupted."
        ), None

    # Legacy DOC
    if filename_lower.endswith(".doc"):

        # Legacy Microsoft Word files normally start with OLE signature
        if file_data.startswith(b"\xD0\xCF\x11\xE0"):
            return True, "", "doc"

        return False, (
            "The uploaded DOC file appears to be invalid or corrupted."
        ), None

    # -----------------------------
    # 4. Unsupported extension
    # -----------------------------
    return False, (
        "Unsupported file type. "
        "Please upload a PDF, DOCX, or DOC resume."
    ), None


# ============================================================
# PDF HYPERLINK EXTRACTION
# ============================================================

def _extract_pdf_hyperlinks(file_data: bytes) -> str:

    urls = []

    try:
        reader = PyPDF2.PdfReader(io.BytesIO(file_data))

        for page in reader.pages:

            if "/Annots" not in page:
                continue

            for annot_ref in page["/Annots"]:

                try:
                    annot = annot_ref.get_object()

                    if annot.get("/Subtype") != "/Link":
                        continue

                    action = annot.get("/A", {})
                    uri = action.get("/URI", "")

                    if uri and isinstance(uri, (str, bytes)):

                        if isinstance(uri, bytes):
                            uri = uri.decode(
                                "utf-8",
                                errors="ignore"
                            )

                        uri = uri.strip()

                        if uri.startswith("http"):
                            urls.append(uri)

                except Exception:
                    pass

    except Exception:
        pass

    return "\n".join(urls)


# ============================================================
# PDF EXTRACTION - PDFPLUMBER
# ============================================================

def _extract_pdf_with_pdfplumber(
    file_data: bytes
) -> str:

    text = ""

    with pdfplumber.open(io.BytesIO(file_data)) as pdf:

        for page in pdf.pages:

            page_text = page.extract_text()

            if page_text:
                text += page_text + "\n"

    if not text.strip():

        raise TextExtractionError(
            "pdfplumber extracted no text",
            user_message="No text could be extracted from the PDF."
        )

    # Extract hyperlinks
    hyperlinks = _extract_pdf_hyperlinks(file_data)

    if hyperlinks:
        text = text.strip() + "\n" + hyperlinks

    return text.strip()


# ============================================================
# PDF EXTRACTION - PYPDF2 FALLBACK
# ============================================================

def _extract_pdf_with_pypdf2(
    file_data: bytes
) -> str:

    text = ""

    pdf_reader = PyPDF2.PdfReader(
        io.BytesIO(file_data)
    )

    for page in pdf_reader.pages:

        page_text = page.extract_text()

        if page_text:
            text += page_text + "\n"

    if not text.strip():

        raise TextExtractionError(
            "PyPDF2 extracted no text",
            user_message="No text could be extracted from the PDF."
        )

    # Extract hyperlinks
    hyperlinks = _extract_pdf_hyperlinks(file_data)

    if hyperlinks:
        text = text.strip() + "\n" + hyperlinks

    return text.strip()


# ============================================================
# PDF TEXT EXTRACTION WITH FALLBACK
# ============================================================

def extract_text_from_pdf(
    file_data: bytes
) -> str:

    try:

        result, used_fallback = with_fallback(
            _extract_pdf_with_pdfplumber,
            _extract_pdf_with_pypdf2,
            file_data,
            log_fallback=True
        )

        if used_fallback:

            log_info(
                "PDF extraction succeeded using the PyPDF2 fallback",
                context="resume_parser"
            )

        return result

    except Exception as e:

        log_error(
            e,
            context="extract_text_from_pdf"
        )

        raise FileParsingError(
            "Failed to extract text from PDF using both "
            "pdfplumber and PyPDF2. "
            "The PDF may be corrupted, password-protected, "
            "or contain only scanned images. "
            "Please ensure it contains selectable text."
        ) from e


# ============================================================
# DOCX TEXT EXTRACTION
# ============================================================

def extract_text_from_docx(
    file_data: bytes
) -> str:

    try:

        doc = Document(
            io.BytesIO(file_data)
        )

        text_parts = []

        # -----------------------------
        # Paragraphs
        # -----------------------------
        for paragraph in doc.paragraphs:

            if paragraph.text.strip():
                text_parts.append(
                    paragraph.text
                )

        # -----------------------------
        # Tables
        # -----------------------------
        for table in doc.tables:

            for row in table.rows:

                for cell in row.cells:

                    if cell.text.strip():
                        text_parts.append(
                            cell.text
                        )

        text = "\n".join(text_parts)

        if not text.strip():

            raise FileParsingError(
                "No text could be extracted from the document. "
                "The document may be empty or corrupted."
            )

        # -----------------------------
        # Extract hyperlinks
        # -----------------------------
        try:

            for rel in doc.part.rels.values():

                if "hyperlink" in rel.reltype.lower():

                    url = rel._target

                    if (
                        isinstance(url, str)
                        and url.startswith("http")
                    ):
                        text += "\n" + url

        except Exception:
            pass

        log_info(
            f"Extracted {len(text)} chars from DOCX",
            context="resume_parser"
        )

        return text.strip()

    except FileParsingError:
        raise

    except Exception as e:

        log_error(
            e,
            context="extract_text_from_docx"
        )

        raise FileParsingError(
            "Failed to extract text from DOCX. "
            "The document may be corrupted or unsupported. "
            "Please try re-saving or converting to PDF."
        ) from e


# ============================================================
# DOC EXTRACTION
# ============================================================

def extract_text_from_doc(
    file_data: bytes
) -> str:

    raise FileParsingError(
        "Legacy .doc format is not supported. "
        "Please convert your resume to .docx or .pdf "
        "and try again."
    )


# ============================================================
# GENERAL TEXT EXTRACTION
# ============================================================

def extract_text(
    file_data: bytes,
    file_type: str
) -> str:

    if file_type == "pdf":

        return extract_text_from_pdf(
            file_data
        )

    elif file_type == "docx":

        return extract_text_from_docx(
            file_data
        )

    elif file_type == "doc":

        return extract_text_from_doc(
            file_data
        )

    else:

        raise FileValidationError(
            f"Invalid file type: {file_type}. "
            "Supported types are: pdf, docx and doc."
        )


# ============================================================
# MAIN RESUME PARSER
# ============================================================

def parse_resume_file(
    file_data: bytes,
    filename: str
) -> Tuple[str, dict]:

    log_info(
        f"Parsing file: {filename}",
        context="parse_resume_file"
    )

    # ========================================================
    # PHASE 01: VALIDATE FILE
    # ========================================================

    try:

        is_valid, error_msg, file_type = validate_file(
            file_data,
            filename
        )

        if not is_valid:

            log_warning(
                f"Validation failed for file {filename}",
                context="parse_resume_file"
            )

            raise FileValidationError(
                error_msg
            )

    except FileValidationError:
        raise

    except Exception as e:

        log_error(
            e,
            context="parse_resume_file_validation"
        )

        raise FileValidationError(
            "Could not validate the uploaded file. "
            "Please ensure it is a valid PDF or DOCX."
        ) from e

    # ========================================================
    # PHASE 02: EXTRACT TEXT
    # ========================================================

    try:

        text = extract_text(
            file_data,
            file_type
        )

        log_info(
            f"Extracted {len(text)} chars from {filename}",
            context="parse_resume_file"
        )

    except FileParsingError:
        raise

    except Exception as e:

        log_error(
            e,
            context="parse_resume_file_extraction"
        )

        raise FileParsingError(
            "An unexpected error occurred while processing "
            "the file. Please try again or contact support "
            "if the problem persists."
        ) from e

    # ========================================================
    # PHASE 03: METADATA
    # ========================================================

    metadata = {
        "filename": filename,
        "file_type": file_type,
        "file_size_bytes": len(file_data),
        "text_length": len(text),
        "success": True,
    }

    return text, metadata